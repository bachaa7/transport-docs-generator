# -*- coding: utf-8 -*-
"""Генерация файлов: docxtpl (Word), openpyxl (ЗНП Excel), LibreOffice (PDF)."""
import re
import shutil
import subprocess
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils.cell import coordinate_to_tuple
from docxtpl import DocxTemplate, InlineImage
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Mm


def safe_name(name):
    """Имя файла без проблемных символов."""
    return re.sub(r"[^0-9A-Za-zА-Яа-я._-]+", "_", str(name)).strip("_") or "doc"


# ---------- Word через docxtpl ----------
def render_docx(template_path, context, out_path, images=None):
    """images: {placeholder: {"path":.., "width":мм | None, "height":мм | None}} —
    вставляются как InlineImage (подпись/печать). Без images рендерится как раньше."""
    tpl = DocxTemplate(str(template_path))
    ctx = dict(context)
    for key, spec in (images or {}).items():
        kw = {}
        if spec.get("width"):
            kw["width"] = Mm(spec["width"])
        if spec.get("height"):
            kw["height"] = Mm(spec["height"])
        ctx[key] = InlineImage(tpl, str(spec["path"]), **kw)
    tpl.render(ctx)
    tpl.save(str(out_path))
    return out_path


def add_floating_stamp(doc_path, stamp_path, anchor_prefix, width_mm=None, height_mm=None,
                       dx_mm=0, dy_mm=0, behind=False, z=251658240):
    """ПЛАВАЮЩАЯ картинка (печать/подпись) поверх текста; фон PNG прозрачный.
    Якорится к абзацу, начинающемуся с anchor_prefix; dx/dy (мм) — смещение от его верх-лево.
    Размер задаётся ШИРИНОЙ (width_mm) ИЛИ ВЫСОТОЙ (height_mm).
    z — слой наложения (relativeHeight): больше = выше (поверх). Печати дают меньше, подписи больше."""
    doc = Document(str(doc_path))
    target = next((p for p in doc.paragraphs if p.text.strip().startswith(anchor_prefix)), None)
    if target is None:
        target = doc.paragraphs[-1]
    run = target.add_run()
    if height_mm:
        run.add_picture(str(stamp_path), height=Mm(height_mm))
    else:
        run.add_picture(str(stamp_path), width=Mm(width_mm or 40))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    cx = inline.find(qn("wp:extent")).get("cx")
    cy = inline.find(qn("wp:extent")).get("cy")
    docpr = inline.find(qn("wp:docPr"))
    pid, pname = docpr.get("id"), docpr.get("name")
    graphic = inline.find(qn("a:graphic"))
    dx, dy = int(Mm(dx_mm)), int(Mm(dy_mm))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "r", "pic")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="{z}" behindDoc="{1 if behind else 0}" locked="0" '
        f'layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:posOffset>{dx}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{dy}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="{pid}" name="{pname}"/><wp:cNvGraphicFramePr/>'
        f'</wp:anchor>')
    anchor.append(graphic)
    drawing.remove(inline)
    drawing.append(anchor)
    doc.save(str(doc_path))
    return doc_path


# ---------- ЗНП через openpyxl по «карте полей» ----------
def _build_lookup(ws):
    tl = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                tl[(r, c)] = (rng.min_row, rng.min_col)
    return tl


def _label(ws, r, tl):
    for col in (1, 2, 3, 4):
        t = tl.get((r, col), (r, col))
        v = ws.cell(t[0], t[1]).value
        if isinstance(v, str) and v.strip():
            return " ".join(v.split())
    return ""


def _yellows(ws, r, tl, color):
    seen, out = set(), []
    for c in range(1, ws.max_column + 1):
        f = ws.cell(r, c).fill
        if f and f.patternType == "solid" and str(getattr(f.fgColor, "rgb", "") or "").endswith(color):
            t = tl.get((r, c), (r, c))
            if t not in seen:
                seen.add(t)
                out.append(t)
    return out


def _resolve(ws, anchor, pick, color):
    """Найти ячейку: строка по подписи (anchor) + жёлтая ячейка (first/last)."""
    tl = _build_lookup(ws)
    for r in range(1, ws.max_row + 1):
        if anchor.lower() in _label(ws, r, tl).lower():
            ys = _yellows(ws, r, tl, color)
            if ys:
                return ys[0] if pick == "first" else ys[-1]
    return None


class _NS:
    def __init__(self, d):
        self.__dict__.update(d)


def _value(istochnik, data):
    """Значение поля: 'const:...' -> литерал; иначе выражение по data (voditel.*, reys.*...)."""
    if istochnik.startswith("const:"):
        return istochnik[len("const:"):]
    ns = {k: _NS(v) for k, v in data.items()}
    return eval(istochnik, {"__builtins__": {}}, ns)


def fill_znp(src_xlsm, karta, data, out_path, images=None):
    """Заполнить .xlsm Альфа по карте полей. Макросы сохраняются (keep_vba).
    images: [{"img":путь, "anchor":"L26", "px":(ширина,высота)}] — подпись/печать (порядок = слои).
    Возвращает (путь, список_ненайденных_полей)."""
    wb = openpyxl.load_workbook(str(src_xlsm), keep_vba=True)
    ws = wb[karta["list_formy"]]
    color = karta["zheltyy_cvet"]
    missed = []
    for pole in karta["polya"]:
        rc = _resolve(ws, pole["anchor"], pole["pick"], color)
        if rc is None:
            missed.append(pole["id"])
            continue
        ws.cell(rc[0], rc[1]).value = _value(pole["istochnik"], data)
    for it in (images or []):                       # печать/подпись поверх области подписи
        p = it["img"]
        if Path(p).exists():
            im = XLImage(str(p))
            nat_w, nat_h = im.width, im.height
            if it.get("h"):                         # высота px, ширина пропорционально (БЕЗ искажения)
                w_px, h_px = int(round(nat_w * it["h"] / nat_h)), it["h"]
            else:
                w_px, h_px = it.get("px", (nat_w, nat_h))
            row, col = coordinate_to_tuple(it["anchor"])   # смещение dx/dy (px) внутри ячейки
            frm = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(it.get("dx", 0)),
                               row=row - 1, rowOff=pixels_to_EMU(it.get("dy", 0)))
            im.anchor = OneCellAnchor(_from=frm, ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(h_px)))
            ws.add_image(im)
    wb.save(str(out_path))
    return out_path, missed


# ---------- PDF через LibreOffice (опционально, мягкая деградация) ----------
def _find_soffice():
    for cand in ("soffice", "libreoffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if shutil.which(cand) or Path(cand).exists():
            return cand
    return None


def to_pdf(path):
    """Конвертация в PDF. Если LibreOffice не установлен — вернёт None (без падения)."""
    soffice = _find_soffice()
    if not soffice:
        return None
    out_dir = Path(path).parent
    subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                    "--outdir", str(out_dir), str(path)],
                   check=True, timeout=120,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = Path(path).with_suffix(".pdf")
    return pdf if pdf.exists() else None
